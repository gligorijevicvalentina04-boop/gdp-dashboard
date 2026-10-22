"""Safe supplementary FreezerFlow features; no destructive schema migrations."""
import io
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from freezerflow.database import connect, now
from freezerflow.services import get_object, children, free_rack_slots, create_storage, log

KINDS={"FREEZER":"SHELF","SHELF":"RACK","RACK":"BOX"}

def styled_workbook(payload):
    wb=load_workbook(io.BytesIO(payload))
    for ws in wb:
        if ws.title=="Instructions":continue
        ws.freeze_panes="A2"
        ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:
            cell.fill=PatternFill('solid',fgColor='12304A')
            cell.font=Font(name='Aptos',size=11,bold=True,color='FFFFFF')
            cell.alignment=Alignment(wrap_text=True)
        for col in ws.columns:
            cells=list(col)
            letter=cells[0].column_letter
            width=min(40,max(15,max(len(str(c.value or '')) for c in cells[:120])+3))
            ws.column_dimensions[letter].width=width
            for i,c in enumerate(cells[1:],2):
                c.font=Font(name='Aptos',size=10,color='12304A')
                if i%2==0:c.fill=PatternFill('solid',fgColor='EAF5F6')
        headers=[c.value for c in ws[1]]
        if 'Action' in headers:
            col=headers.index('Action')+1
            from openpyxl.utils import get_column_letter
            letter=get_column_letter(col)
            dv=DataValidation(type='list',formula1='"ADD,UPDATE,REMOVE"')
            ws.add_data_validation(dv)
            dv.add(f'{letter}2:{letter}1000')
    out=io.BytesIO();wb.save(out);return out.getvalue()

def freezer_shelves_template(uid):
    from openpyxl import Workbook
    wb=Workbook();ws=wb.active;ws.title='Etageres'
    ws.append(['Nom étagère','Étagère UID','Action'])
    for item in children(uid,'SHELF'):ws.append([item['name'],item['uid'],'UPDATE'])
    ws.append(['','','ADD'])
    out=io.BytesIO();wb.save(out);return styled_workbook(out.getvalue())

def apply_shelves(uid,rows,user=''):
    parent=get_object(uid)
    if not parent or parent['kind']!='FREEZER':return 0,['Congélateur introuvable']
    count=0;errors=[]
    for line,r in enumerate(rows,2):
        try:
            action=str(r.get('Action') or 'ADD').strip().upper()
            name=str(r.get('Nom étagère') or '').strip()
            objuid=str(r.get('Étagère UID') or '').strip()
            if not name and not objuid:continue
            if action not in ('ADD','UPDATE','REMOVE'):raise ValueError('Action inconnue')
            with connect() as c:
                existing=c.execute('SELECT * FROM storage_objects WHERE uid=? AND active=1',(objuid,)).fetchone() if objuid else None
                if action!='ADD' and (not existing or existing['kind']!='SHELF' or existing['parent_uid']!=uid):raise ValueError('Étagère UID non trouvée dans ce congélateur')
                if action=='ADD':
                    if objuid:raise ValueError('Ne pas fournir un UID pour ADD')
                    if not name:raise ValueError('Nom requis')
                    create_storage('SHELF',name,uid)
                elif action=='UPDATE':
                    if not name:raise ValueError('Nom requis')
                    c.execute('UPDATE storage_objects SET name=? WHERE uid=?',(name,objuid))
                    # Journaliser après validation de la transaction
                else:
                    if c.execute('SELECT 1 FROM storage_objects WHERE parent_uid=? AND active=1',(objuid,)).fetchone():raise ValueError('Archiver les enfants avant cette étagère')
                    c.execute('UPDATE storage_objects SET active=0 WHERE uid=?',(objuid,))
                    # Journaliser après validation de la transaction
            if action in ('UPDATE','REMOVE'):
                log(objuid,'MODIFICATION' if action=='UPDATE' else 'ARCHIVAGE',user=user,details='Action via Excel')
            count+=1
        except Exception as e:errors.append(f'Ligne {line} : {e}')
    return count,errors

def create_batch(kind,parent_uid,quantity,base_name,rows=9,cols=9,rack_rows=3,rack_cols=3):
    parent=get_object(parent_uid)
    if not parent or KINDS.get(parent['kind'])!=kind:raise ValueError('Niveau parent incorrect')
    quantity=int(quantity)
    if not 1<=quantity<=100:raise ValueError('Nombre entre 1 et 100')
    if kind=='BOX' and quantity>len(free_rack_slots(parent_uid)):raise ValueError('Pas assez de places libres dans ce rack')
    created=[]
    for i in range(quantity):
        name=f'{base_name.strip()} {i+1}' if quantity>1 else base_name.strip()
        if not name:raise ValueError('Nom requis')
        if kind=='BOX':uid=create_storage(kind,name,parent_uid,rows,cols,rack_position=free_rack_slots(parent_uid)[0])
        elif kind=='RACK':uid=create_storage(kind,name,parent_uid,capacity=rack_rows*rack_cols,rack_rows=rack_rows,rack_cols=rack_cols)
        else:uid=create_storage(kind,name,parent_uid)
        created.append(uid)
    return created
