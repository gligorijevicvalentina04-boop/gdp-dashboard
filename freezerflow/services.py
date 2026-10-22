import io,uuid,re,os
from urllib.parse import quote, unquote
from pathlib import Path
import qrcode
from PIL import Image,ImageOps
from openpyxl import Workbook,load_workbook
from freezerflow.database import execute,fetchall,fetchone,now,connect

PREFIX={"FREEZER":"FZ","SHELF":"SH","RACK":"RK","BOX":"BOX","TUBE":"T"}
PHOTO_DIR=Path(__file__).resolve().parent.parent / "data" / "photos";PHOTO_DIR.mkdir(parents=True,exist_ok=True)

def new_uid(kind):return f"{PREFIX[kind]}-{uuid.uuid4().hex.upper()}"

def create_storage(kind,name,parent_uid=None,rows=None,cols=None,capacity=None,rack_rows=None,rack_cols=None,rack_position=None):
    uid=new_uid(kind)
    execute("""INSERT INTO storage_objects(uid,kind,name,parent_uid,rows_count,cols_count,capacity,rack_rows,rack_cols,rack_position,created_at)
    VALUES(?,?,?,?,?,?,?,?,?,?,?)""",(uid,kind,name.strip(),parent_uid,rows,cols,capacity,rack_rows,rack_cols,rack_position,now()))
    log(uid,"CREATION","",parent_uid,"",f"Création {kind}: {name}");return uid

def roots():return fetchall("SELECT * FROM storage_objects WHERE kind='FREEZER' AND active=1 ORDER BY name")
def children(uid,kind=None):
    if kind:return fetchall("SELECT * FROM storage_objects WHERE parent_uid=? AND kind=? AND active=1 ORDER BY name",(uid,kind))
    return fetchall("SELECT * FROM storage_objects WHERE parent_uid=? AND active=1 ORDER BY kind,name",(uid,))
def get_object(uid):return fetchone("SELECT * FROM storage_objects WHERE uid=? AND active=1",(uid,))
def all_boxes():return fetchall("SELECT * FROM storage_objects WHERE kind='BOX' AND active=1 ORDER BY name")
def all_racks():return fetchall("SELECT * FROM storage_objects WHERE kind='RACK' AND active=1 ORDER BY name")
def tube(uid):return fetchone("SELECT * FROM tubes WHERE uid=? AND active=1",(uid,))
def tubes_in_box(uid):return fetchall("SELECT * FROM tubes WHERE box_uid=? AND active=1 AND status='PRESENT' ORDER BY position",(uid,))
def breadcrumb(uid):
    a=[];o=get_object(uid)
    while o:a.append(o);o=get_object(o["parent_uid"]) if o["parent_uid"] else None
    return list(reversed(a))
def log(uid,action,fr="",to="",user="",details=""):
    execute("INSERT INTO movements(object_uid,action,from_location,to_location,user_name,details,created_at) VALUES(?,?,?,?,?,?,?)",
            (uid,action,fr,to,user,details,now()))
def rename_object(uid,name,user=""):
    o=get_object(uid)
    if not o or not name.strip():return False
    execute("UPDATE storage_objects SET name=? WHERE uid=?",(name.strip(),uid))
    log(uid,"MODIFICATION","","",user,f"Nom: {o['name']} → {name.strip()}");return True
def create_tube(uid,name,box,pos,user=""):
    execute("INSERT INTO tubes(uid,sample_name,box_uid,position,status,created_at) VALUES(?,?,?,?, 'PRESENT',?)",
            (uid.strip(),name.strip(),box,pos,now()));log(uid,"AJOUT","",f"{box}/{pos}",user,"Tube ajouté")

def create_multiple_tubes(sample_name,box_uid,quantity,user="",prefix="T",volume="",unit="mL",lot=""):
    try:quantity=int(quantity)
    except (ValueError,TypeError):return [],["Quantité invalide."]
    if not sample_name.strip():return [],["Nom d'échantillon obligatoire."]
    free=free_positions(box_uid)
    if quantity<1 or quantity>len(free):return [],[f"Positions libres : {len(free)}."]
    try:vol=float(volume) if volume not in ("",None) else None
    except (ValueError,TypeError):return [],["Volume invalide."]
    if vol is not None and vol<0:return [],["Volume négatif interdit."]
    created=[]
    try:
        with connect() as c:
            for pos in free[:quantity]:
                uid=new_uid("TUBE")
                c.execute("""INSERT INTO tubes(uid,sample_name,box_uid,position,status,created_at,volume,volume_unit,lot)
                             VALUES(?,?,?,?, 'PRESENT',?,?,?,?)""",
                          (uid,sample_name.strip(),box_uid,pos,now(),vol,unit,lot))
                c.execute("""INSERT INTO movements(object_uid,action,from_location,to_location,user_name,details,created_at)
                             VALUES(?,?,?,?,?,?,?)""",(uid,"AJOUT","",f"{box_uid}/{pos}",user,"Ajout multiple",now()))
                created.append({"Tube UID":uid,"Échantillon":sample_name.strip(),"Position":pos,"Volume":vol,"Unité":unit,"Lot":lot})
        return created,[]
    except Exception as e:return [],[str(e)]

def update_tube(uid,name,user=""):
    t=tube(uid)
    if not t:return False
    execute("UPDATE tubes SET sample_name=? WHERE uid=?",(name.strip(),uid))
    log(uid,"MODIFICATION","","",user,f"Échantillon: {t['sample_name'] or ''} → {name.strip()}");return True
def remove_tube(uid,user=""):
    t=tube(uid)
    if not t:return False
    execute("UPDATE tubes SET status='RETIRED',removed_at=? WHERE uid=?",(now(),uid))
    log(uid,"RETRAIT",f"{t['box_uid']}/{t['position']}","",user,"Tube retiré");return True
def restore_tube(uid,box,pos,user=""):
    execute("UPDATE tubes SET box_uid=?,position=?,status='PRESENT',removed_at=NULL WHERE uid=?",(box,pos,uid))
    log(uid,"REMISE","",f"{box}/{pos}",user,"Tube remis")
def move_tube(uid,box,pos,user=""):
    t=tube(uid)
    if not t:return False
    old=f"{t['box_uid']}/{t['position']}"
    execute("UPDATE tubes SET box_uid=?,position=? WHERE uid=?",(box,pos,uid))
    log(uid,"DEPLACEMENT",old,f"{box}/{pos}",user,"Tube déplacé");return True
def archive_object(uid,user=""):
    o=get_object(uid)
    if not o:return False,"Objet introuvable."
    if fetchone("SELECT COUNT(*) n FROM storage_objects WHERE parent_uid=? AND active=1",(uid,))["n"]:
        return False,"Supprimez/archivez d'abord les éléments contenus."
    if o["kind"]=="BOX" and fetchone("SELECT COUNT(*) n FROM tubes WHERE box_uid=? AND active=1 AND status='PRESENT'",(uid,))["n"]:
        return False,"Retirez d'abord les tubes présents dans cette boîte."
    execute("UPDATE storage_objects SET active=0 WHERE uid=?",(uid,));log(uid,"ARCHIVAGE","","",user,"Suppression logique");return True,""
def object_url(uid):
    base=os.getenv("FREEZERFLOW_PUBLIC_URL", "").strip().rstrip("/")
    return f"{base}/?item={quote(uid)}" if base else ""

def qr_png(v):
    im=qrcode.make(v);b=io.BytesIO();im.save(b,format="PNG");return b.getvalue()
def free_positions(uid):
    b=get_object(uid)
    if not b or not b["rows_count"] or not b["cols_count"]:return []
    used={t["position"] for t in tubes_in_box(uid)}
    return [f"{chr(65+r)}{c}" for r in range(b["rows_count"]) for c in range(1,b["cols_count"]+1) if f"{chr(65+r)}{c}" not in used]
def rack_slots(uid):
    r=get_object(uid);rows=(r or {}).get("rack_rows") or 3;cols=(r or {}).get("rack_cols") or 3
    return [f"{chr(65+i)}{j}" for i in range(rows) for j in range(1,cols+1)]
def free_rack_slots(uid):
    used={x["rack_position"] for x in children(uid,"BOX") if x.get("rack_position")}
    return [x for x in rack_slots(uid) if x not in used]
def set_box_photo(uid,f):
    ext=Path(getattr(f,"name","photo.jpg")).suffix.lower()
    if ext not in {".jpg",".jpeg",".png"}:ext=".jpg"
    p=PHOTO_DIR/f"{uid}{ext}";p.write_bytes(f.getvalue());execute("UPDATE storage_objects SET photo_path=? WHERE uid=?",(str(p),uid));log(uid,"PHOTO",details="Photo ajoutée/remplacée")
def first_free_location():
    for b in all_boxes():
        f=free_positions(b["uid"])
        if f:return b,f[0]
    return None,None

def descendants(uid):
    out=[]
    for c in children(uid):
        out.append(c);out.extend(descendants(c["uid"]))
    return out

def export_excel(uid):
    o=get_object(uid);bio=io.BytesIO();wb=Workbook();ws=wb.active;ws.title="Inventaire"
    headers=["Niveau","Nom","UID","Parent UID","Position rack","Position tube","Tube UID","Échantillon","Statut"]
    ws.append(headers)
    if not o:raise ValueError("Objet introuvable.")
    objs=[o]+descendants(uid)
    for x in objs:
        if x["kind"]=="BOX":
            ts=tubes_in_box(x["uid"])
            if ts:
                for t in ts:ws.append([x["kind"],x["name"],x["uid"],x["parent_uid"],x.get("rack_position"),t["position"],t["uid"],t["sample_name"],t["status"]])
            else:ws.append([x["kind"],x["name"],x["uid"],x["parent_uid"],x.get("rack_position"),"","","",""])
        else:ws.append([x["kind"],x["name"],x["uid"],x["parent_uid"],x.get("rack_position"),"","","",""])
    ws.freeze_panes="A2";wb.save(bio);return bio.getvalue()

def box_template(uid):
    b=get_object(uid);bio=io.BytesIO();wb=Workbook();ws=wb.active;ws.title="Tubes"
    ws.append(["Position","Tube UID","Échantillon","Volume","Unité","Lot","Action"])
    present={x["position"]:x for x in tubes_in_box(uid)}
    for p in [f"{chr(65+r)}{c}" for r in range(b["rows_count"] or 0) for c in range(1,(b["cols_count"] or 0)+1)]:
        x=present.get(p);ws.append([p,x["uid"] if x else "",x["sample_name"] if x else "",x.get("volume") if x else "",x.get("volume_unit") if x else "",x.get("lot") if x else "", "UPDATE" if x else "ADD"])
    info=wb.create_sheet("Instructions");info["A1"]="ADD: ajouter, UPDATE: modifier par UID, REMOVE: retirer. Une ligne absente ne supprime rien."
    wb.save(bio);return bio.getvalue()

def rack_template(uid):
    bio=io.BytesIO();wb=Workbook();ws=wb.active;ws.title="Boites";ws.append(["Position rack","Nom boîte","Lignes","Colonnes","Boîte UID","Action"])
    existing={x["rack_position"]:x for x in children(uid,"BOX")}
    for p in rack_slots(uid):
        x=existing.get(p);ws.append([p,x["name"] if x else "",x["rows_count"] if x else 9,x["cols_count"] if x else 9,x["uid"] if x else "","UPDATE" if x else "ADD"])
    wb.save(bio);return bio.getvalue()

def shelf_template(uid):
    bio=io.BytesIO();wb=Workbook();ws=wb.active;ws.title="Racks";ws.append(["Nom rack","Niveaux","Boîtes par niveau","Rack UID","Action"])
    for x in children(uid,"RACK"):ws.append([x["name"],x.get("rack_rows") or 3,x.get("rack_cols") or 3,x["uid"],"UPDATE"])
    ws.append(["",3,3,"","ADD"])
    wb.save(bio);return bio.getvalue()

def parse_excel(upload,kind):
    wb=load_workbook(io.BytesIO(upload.getvalue()),data_only=True);ws=wb.active;rows=list(ws.iter_rows(values_only=True))
    if not rows:return []
    headers=[str(x or "").strip() for x in rows[0]];out=[]
    for vals in rows[1:]:
        d={headers[i]:vals[i] if i<len(vals) else None for i in range(len(headers))}
        if any(v not in (None,"") for v in d.values()):out.append(d)
    return out

def import_rows(parent_uid,kind,rows,user=""):
    """Apply Excel changes by UID; never silently delete absent rows."""
    added=0;errors=[]
    for i,r in enumerate(rows,2):
        try:
            action=str(r.get("Action") or "").strip().upper()
            if action not in ("", "ADD", "UPDATE", "REMOVE"):
                raise ValueError("Action invalide (ADD, UPDATE ou REMOVE)")
            if kind=="BOX":
                pos=str(r.get("Position") or "").strip().upper()
                uid=str(r.get("Tube UID") or "").strip()
                name=str(r.get("Échantillon") or "").strip()
                if not uid and not name:continue
                existing=tube(uid) if uid else None
                if action=="REMOVE":
                    if not existing or existing["box_uid"]!=parent_uid:raise ValueError("Tube non trouvé dans cette boîte")
                    remove_tube(uid,user);added+=1;continue
                if existing:
                    if existing["box_uid"]!=parent_uid:raise ValueError("UID associé à une autre boîte")
                    if pos and pos!=existing["position"]:
                        if pos not in free_positions(parent_uid):raise ValueError(f"position {pos} indisponible")
                        move_tube(uid,parent_uid,pos,user)
                    update_tube(uid,name,user)
                else:
                    if action=="UPDATE":raise ValueError("UID à modifier introuvable")
                    if pos not in free_positions(parent_uid):raise ValueError(f"position {pos} indisponible")
                    create_tube(uid or new_uid("TUBE"),name,parent_uid,pos,user)
            elif kind=="RACK":
                pos=str(r.get("Position rack") or "").strip().upper()
                name=str(r.get("Nom boîte") or "").strip()
                uid=str(r.get("Boîte UID") or "").strip()
                existing=get_object(uid) if uid else None
                if action=="REMOVE":
                    if not existing or existing["parent_uid"]!=parent_uid:raise ValueError("Boîte non trouvée dans ce rack")
                    ok,msg=archive_object(uid,user)
                    if not ok:raise ValueError(msg)
                elif existing:
                    if existing["kind"]!="BOX" or existing["parent_uid"]!=parent_uid:raise ValueError("UID de boîte incorrect")
                    if pos and pos!=existing["rack_position"]:
                        if pos not in free_rack_slots(parent_uid):raise ValueError("Position rack occupée")
                        execute("UPDATE storage_objects SET rack_position=? WHERE uid=?",(pos,uid))
                    if name:rename_object(uid,name,user)
                else:
                    if not name:continue
                    if action=="UPDATE":raise ValueError("UID à modifier introuvable")
                    if pos not in free_rack_slots(parent_uid):raise ValueError(f"emplacement {pos} indisponible")
                    create_storage("BOX",name,parent_uid,int(r.get("Lignes") or 9),int(r.get("Colonnes") or 9),rack_position=pos)
            elif kind=="SHELF":
                name=str(r.get("Nom rack") or "").strip()
                uid=str(r.get("Rack UID") or "").strip()
                existing=get_object(uid) if uid else None
                if action=="REMOVE":
                    if not existing or existing["parent_uid"]!=parent_uid:raise ValueError("Rack non trouvé")
                    ok,msg=archive_object(uid,user)
                    if not ok:raise ValueError(msg)
                elif existing:
                    if existing["kind"]!="RACK" or existing["parent_uid"]!=parent_uid:raise ValueError("UID rack incorrect")
                    if name:rename_object(uid,name,user)
                else:
                    if not name:continue
                    if action=="UPDATE":raise ValueError("UID à modifier introuvable")
                    rr=int(r.get("Niveaux") or 3);rc=int(r.get("Boîtes par niveau") or 3)
                    create_storage("RACK",name,parent_uid,capacity=rr*rc,rack_rows=rr,rack_cols=rc)
            added+=1
        except Exception as e:errors.append(f"Ligne {i}: {e}")
    return added,errors

def freezer_template():
    bio=io.BytesIO();wb=Workbook();ws=wb.active;ws.title="Congélateurs"
    ws.append(["Nom", "Type", "Action"]);ws.append(["Congélateur test","Congélateur", "ADD"])
    wb.save(bio);return bio.getvalue()

def import_freezers(rows,user=""):
    done=0;errors=[]
    for i,r in enumerate(rows,2):
        try:
            name=str(r.get("Nom") or "").strip();action=str(r.get("Action") or "ADD").strip().upper()
            if not name:continue
            if action!="ADD":raise ValueError("Le modèle global accepte uniquement ADD pour éviter une suppression involontaire")
            create_storage("FREEZER",name);done+=1
        except Exception as e:errors.append(f"Ligne {i}: {e}")
    return done,errors

def ocr_image(upload):
    try:
        import pytesseract
        img=ImageOps.exif_transpose(Image.open(io.BytesIO(upload.getvalue())))
        text=pytesseract.image_to_string(img)
        found=[]
        # Expected readable patterns: A1 T-001 SampleName, one tube per line.
        for line in text.splitlines():
            m=re.match(r"^\s*([A-Z][0-9]{1,2})\s+([A-Za-z0-9_.-]+)(?:\s+(.*))?$",line.strip())
            if m:found.append({"Position":m.group(1),"Tube UID":m.group(2),"Échantillon":m.group(3) or ""})
        return text,found,None
    except Exception as e:return "",[],str(e)
